"""Загрузка товаров из таблицы: шаблон программы (.xlsx) и выгрузка в CSV, как из Hoog."""

from datetime import datetime

import pytest
from openpyxl import Workbook

from app import stock_import
from app.db import Database, UserError
from app.stock import Product

NOW = datetime(2026, 10, 29, 12, 0)


@pytest.fixture
def db():
    database = Database(":memory:")
    yield database
    database.close()


def test_number():
    assert stock_import.number("1 250,50 ₽") == 1250.5
    assert stock_import.number("1.250.50") == 1250.5
    assert stock_import.number(98) == 98 and stock_import.number("") is None and stock_import.number(None) is None


def test_template_xlsx(db, tmp_path):
    book = Workbook()
    sheet = book.active
    sheet.title = "Товары"
    sheet.append(["Название", "Раздел", "Единица", "Цена продажи, ₽", "Закупочная цена, ₽", "Остаток",
                  "Мало, если остаток ≤"])
    sheet.append(["Роза Кения", "цветы", "шт", 250, 85, 120, 30])
    sheet.append(["Лента атласная", "Упаковка", "м", 50, 8, 100.5, None])
    sheet.append([None, None, None, None, None, None, None])
    sheet.append(["Анкерная лента", "Упаковка", "м", 0, 5, None, 10])
    book.create_sheet("Как заполнять").append(["Пояснения"])
    path = str(tmp_path / "товары.xlsx")
    book.save(path)
    rows, problems = stock_import.parse(db, stock_import.read_table(path))
    assert problems == [] and [r.action for r in rows] == ["добавить"] * 3
    assert (rows[0].category, rows[0].price, rows[0].cost, rows[0].stock, rows[0].min_stock) == ("Цветы", 250, 85, 120, 30)
    assert stock_import.apply(db, rows, None, NOW) == (3, 0)
    rose = next(p for p in db.products() if p.name == "Роза Кения")
    assert (rose.stock, rose.avg_cost, rose.min_stock, rose.price) == (120, 85, 30, 250)
    ribbon = next(p for p in db.products() if p.name == "Лента атласная")
    assert (ribbon.unit, ribbon.stock, ribbon.category) == ("м", 100.5, "Упаковка")
    assert db.supplies()[0].supplier == "Загрузка из таблицы"


def test_hoog_like_csv_updates_existing(db, tmp_path):
    db.save_product(Product(0, "Гвоздика красная", "Цветы", price=90))
    path = tmp_path / "hoog.csv"
    path.write_text("Выгрузка остатков\n\nНаименование;Категория;Кол-во;Цена;Себестоимость\n"
                    "Гвоздика красная;Штучный;80;98;35\nЭксплоэр;Штучный;45;550;180\n"
                    "Эксплоэр;Штучный;40;550;180\nПлохая;Штучный;1;-5;1\n", encoding="cp1251")
    rows, problems = stock_import.parse(db, stock_import.read_table(str(path)))
    assert [(r.name, r.action) for r in rows] == [("Гвоздика красная", "обновить"), ("Эксплоэр", "добавить")]
    assert any("уже был выше" in p for p in problems) and any("отрицательная" in p for p in problems)
    assert stock_import.apply(db, rows, None, NOW) == (1, 1)
    carnation = next(p for p in db.products() if p.name == "Гвоздика красная")
    assert (carnation.price, carnation.stock, carnation.avg_cost, carnation.category) == (98, 80, 35, "Штучный")
    assert next(p for p in db.products() if p.name == "Эксплоэр").stock == 40


def test_bad_files(db, tmp_path):
    (tmp_path / "x.csv").write_text("a;b\n1;2\n", encoding="utf-8")
    rows, problems = stock_import.parse(db, stock_import.read_table(str(tmp_path / "x.csv")))
    assert rows == [] and "Не нашёл столбец" in problems[0]
    (tmp_path / "old.xls").write_bytes(b"\xd0\xcf")
    with pytest.raises(UserError, match=".xlsx"):
        stock_import.read_table(str(tmp_path / "old.xls"))
    (tmp_path / "broken.xlsx").write_bytes(b"not a zip")
    with pytest.raises(UserError, match="Не получается открыть"):
        stock_import.read_table(str(tmp_path / "broken.xlsx"))


def test_hoog_stock_export(db, tmp_path):
    """Выгрузка остатков Hoog: ID, Магазин, «Товар (номер)», Количество строкой; цен и разделов нет."""
    book = Workbook()
    sheet = book.active
    sheet.append(["ID\n(Click to sort ascending)", "Магазин\n(Click to clear sorting)", "Товар\n(Click to clear sorting)",
                  "Количество\n(Click to clear sorting)", None])
    for row in ([1, "Страна Чудес (6827)", "Кения (565001)", "301.00"],
                [2, "Страна Чудес (6827)", "Матиола (565285)", "30.00"],
                [3, "Страна Чудес (6827)", "Матиола (380111)", "2.00"],  # то же имя, другой товар Hoog — сложить
                [4, "Страна Чудес (6827)", "Букет 20 (567164)", "1.00"],  # собранный в Hoog — не товар
                [5, "Страна Чудес (6827)", "Пинк Лола (380394)", "0.00"],  # закончился
                [6, "Страна Чудес (6827)", "Фольгированные Шары н1 (400001)", "2.00"],
                [7, "Страна Чудес (6827)", "Банты н3 (400002)", "40.00"],
                [8, "Страна Чудес (6827)", "Мыльный бокс н301 (400003)", "1.00"],
                [9, "Страна Чудес (6827)", "Эвкалипт (400004)", "9.00"],
                [10, "Страна Чудес (6827)", "Гипсофила (400005)", "17.60"]):
        sheet.append(row)
    path = str(tmp_path / "Остатки на складах.xlsx")
    book.save(path)
    rows, problems = stock_import.parse(db, stock_import.read_table(path))
    got = {r.name: (r.category, r.stock) for r in rows}
    assert got == {"Кения": ("Цветы", 301), "Матиола": ("Цветы", 32), "Фольгированные Шары н1": ("Сопутствующие", 2),
                   "Банты н3": ("Упаковка", 40), "Мыльный бокс н301": ("Сопутствующие", 1),
                   "Эвкалипт": ("Зелень", 9), "Гипсофила": ("Цветы", 17.6)}
    text = " ".join(problems)
    assert "нулевым остатком — 1" in text and "Собранных букетов из Hoog" in text and "— 1: остатки сложил" in text
    assert "Цен в таблице нет" in text
    assert stock_import.apply(db, rows, None, NOW) == (7, 0)
    stock_flower = next(p for p in db.products() if p.name == "Матиола")
    stock_flower.category = "Сухоцветы"  # поправили раздел руками
    db.save_product(stock_flower)
    # следом — прайс с ценами: цены обновятся, остатки и поправленный раздел не тронутся
    prices = tmp_path / "цены.csv"
    prices.write_text("Товар;Цена\nКения (565001);250\nМатиола;180\n", encoding="utf-8")
    rows, _ = stock_import.parse(db, stock_import.read_table(str(prices)))
    assert stock_import.apply(db, rows, None, NOW) == (0, 2)
    kenya = next(p for p in db.products() if p.name == "Кения")
    assert (kenya.price, kenya.stock, kenya.category) == (250, 301, "Цветы")
    assert next(p for p in db.products() if p.name == "Матиола").category == "Сухоцветы"
