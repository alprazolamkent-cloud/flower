"""Загрузка товаров из таблицы (Excel .xlsx или .csv): шаблон программы или выгрузка из Hoog.

Столбцы узнаются по названию («Название» / «Наименование», «Цена», «Остаток»…), порядок не важен.
Новые товары добавляются, с тем же названием — обновляются; остатки заносятся одной поставкой
«Загрузка из таблицы» (новые) или инвентаризацией (уже были), чтобы себестоимость считалась сразу.
"""

import csv
import re
from dataclasses import dataclass
from datetime import datetime

from .stock import CATEGORIES, Product, StockError

# столбец → слова, по которым его узнать (первое совпадение в заголовке)
COLUMNS = {
    "name": ("название", "наименование", "товар", "номенклатура", "name"),
    "category": ("раздел", "категория", "группа", "тип", "category"),
    "unit": ("единица", "ед.", "ед ", "ед", "unit"),
    "cost": ("закуп", "себестоим", "входная", "cost"),
    "price": ("цена продажи", "розничн", "цена", "стоимость", "price"),
    "min_stock": ("мало", "минимал", "мин.", "min"),  # раньше «остатка»: «Мало, если остаток ≤»
    "stock": ("остаток", "количество", "кол-во", "кол.", "наличие", "stock", "qty"),
}


@dataclass
class ImportRow:
    line: int  # строка в таблице (для понятных ошибок)
    name: str
    category: str | None = None
    unit: str | None = None
    price: int | None = None
    cost: float | None = None
    stock: float | None = None
    min_stock: float | None = None
    existing: Product | None = None

    @property
    def action(self) -> str:
        return "обновить" if self.existing else "добавить"


def number(value) -> float | None:
    """«1 250,50 ₽» → 1250.5; пусто — None."""
    if value is None:
        return None
    if isinstance(value, (int, float)):
        return float(value)
    text = re.sub(r"[^\d,.\-]", "", str(value).replace("\xa0", "").replace(" ", "")).replace(",", ".")
    if text.count(".") > 1:  # 1.250.50 — разделители тысяч
        head, _, tail = text.rpartition(".")
        text = head.replace(".", "") + "." + tail
    try:
        return float(text) if text not in ("", "-", ".") else None
    except ValueError:
        return None


def _header_map(header: list) -> dict[str, int]:
    found: dict[str, int] = {}
    cells = [str(h or "").strip().lower() for h in header]
    for key, words in COLUMNS.items():  # порядок важен: «закупочная цена» — это cost, а не price
        for i, cell in enumerate(cells):
            if i in found.values() or not cell:
                continue
            if any(cell.startswith(w) or f" {w}" in f" {cell}" for w in words):
                found[key] = i
                break
    return found


def read_table(path: str) -> list[list]:
    """Все строки первой (или «Товары») таблицы файла."""
    if path.lower().endswith((".xlsx", ".xlsm")):
        try:
            from openpyxl import load_workbook
        except ImportError:  # pragma: no cover — в сборке есть
            raise StockError("Не получается открыть Excel. Сохраните таблицу как CSV и загрузите её.") from None
        try:
            book = load_workbook(path, read_only=True, data_only=True)
        except Exception:  # noqa: BLE001 — битый или не тот файл
            raise StockError("Не получается открыть файл. Это точно таблица Excel (.xlsx)?") from None
        sheet = book["Товары"] if "Товары" in book.sheetnames else book.worksheets[0]
        rows = [list(r) for r in sheet.iter_rows(values_only=True)]
        book.close()
        return rows
    if path.lower().endswith(".xls"):
        raise StockError("Старый формат .xls не открывается. В Excel: «Сохранить как» → «Книга Excel (.xlsx)».")
    for encoding in ("utf-8-sig", "cp1251"):
        try:
            with open(path, encoding=encoding, newline="") as f:
                text = f.read()
            break
        except UnicodeDecodeError:
            continue
    else:
        raise StockError("Не получается прочитать файл.")
    try:
        delimiter = csv.Sniffer().sniff(text[:4096], delimiters=";,\t").delimiter
    except csv.Error:
        delimiter = ";"  # русский Excel сохраняет CSV через «;»
    return [row for row in csv.reader(text.splitlines(), delimiter=delimiter)]


def parse(db, rows: list[list]) -> tuple[list[ImportRow], list[str]]:
    """Строки таблицы → что добавить / обновить, и список проблем."""
    problems: list[str] = []
    start = next((i for i, r in enumerate(rows[:20]) if "name" in _header_map(r)), None)
    if start is None:
        return [], ["Не нашёл столбец с названием товара («Название» или «Наименование»)."]
    columns = _header_map(rows[start])
    if "price" not in columns and "stock" not in columns:
        problems.append("Нет столбцов «Цена» и «Остаток» — загрузятся только названия.")
    existing = {p.name.casefold(): p for p in db.products()}
    seen: dict[str, ImportRow] = {}

    def cell(row, key):
        i = columns.get(key)
        return row[i] if i is not None and i < len(row) else None

    for line, row in enumerate(rows[start + 1:], start=start + 2):
        name = str(cell(row, "name") or "").strip()
        if not name:
            continue
        category = str(cell(row, "category") or "").strip() or None
        if category:
            category = next((c for c in CATEGORIES if c.casefold() == category.casefold()), category)
        unit = str(cell(row, "unit") or "").strip() or None
        price, cost = number(cell(row, "price")), number(cell(row, "cost"))
        stock, low = number(cell(row, "stock")), number(cell(row, "min_stock"))
        if any(v is not None and v < 0 for v in (price, cost, low)):
            problems.append(f"Строка {line} «{name}»: отрицательная цена — пропустил.")
            continue
        if name.casefold() in seen:
            problems.append(f"Строка {line}: «{name}» уже был выше — взял последнюю строку.")
        seen[name.casefold()] = ImportRow(line, name, category, unit, None if price is None else round(price), cost,
                                          stock, low, existing.get(name.casefold()))
    if not seen:
        problems.append("В таблице нет ни одного товара.")
    return list(seen.values()), problems


def apply(db, rows: list[ImportRow], user_id: int | None, now: datetime) -> tuple[int, int]:
    """Занести в базу. Возвращает (добавлено, обновлено)."""
    added = updated = 0
    supply: list[tuple[int, float, float]] = []
    for r in rows:
        p = r.existing or Product(0, r.name)
        if r.category:
            p.category = r.category
        if r.unit:
            p.unit = r.unit
        if r.price is not None:
            p.price = r.price
        if r.min_stock is not None:
            p.min_stock = r.min_stock
        db.save_product(p)
        if r.existing is None:
            added += 1
            if r.stock:
                supply.append((p.id, r.stock, r.cost or 0))
            continue
        updated += 1
        if r.cost is not None:  # в таблице закупочная цена — ей и верим
            db.conn.execute("UPDATE products SET avg_cost=? WHERE id=?", (r.cost, p.id))
            db.conn.commit()
        if r.stock is not None:
            db.count_stock(p.id, max(r.stock, 0), user_id, now)
    if supply:
        db.add_supply("Загрузка из таблицы", supply, user_id, now, "остатки на начало")
    return added, updated
